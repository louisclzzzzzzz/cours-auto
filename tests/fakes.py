"""Faux services Google Drive et Notion (en mémoire) et petits documents de test (PDF, PPTX)."""

from __future__ import annotations

import itertools
import re
from types import SimpleNamespace

from googleapiclient.errors import HttpError

# --- Drive -----------------------------------------------------------------------------------------


class _Req:
    def __init__(self, fn):
        self._fn = fn

    def execute(self, num_retries: int = 0):
        return self._fn()


class FakeDriveFiles:
    def __init__(self, store: dict, log: list):
        self.store = store
        self.log = log
        self._ids = itertools.count(1)

    def _public(self, f: dict) -> dict:
        return {k: f[k] for k in ("id", "name", "mimeType", "trashed", "webViewLink")}

    def get(self, fileId, fields=None):
        def run():
            f = self.store.get(fileId)
            if not f:
                raise HttpError(SimpleNamespace(status=404, reason="Not Found"), b"not found")
            return self._public(f)
        return _Req(run)

    def create(self, body, media_body=None, fields=None):
        def run():
            fid = f"f{next(self._ids)}"
            mime = body.get("mimeType") or (media_body.mimetype() if media_body else "application/octet-stream")
            f = {
                "id": fid, "name": body["name"], "mimeType": mime, "parents": body.get("parents", []),
                "trashed": False, "webViewLink": f"https://drive.example/{fid}",
                "content": media_body.getbytes(0, media_body.size()) if media_body else None,
                "source_mime": media_body.mimetype() if media_body else None,
            }
            self.store[fid] = f
            self.log.append(("create", fid, body["name"]))
            return self._public(f)
        return _Req(run)

    def update(self, fileId, body=None, media_body=None, fields=None):
        def run():
            f = self.store[fileId]
            if body and "name" in body:
                f["name"] = body["name"]
            if media_body is not None:
                f["content"] = media_body.getbytes(0, media_body.size())
                f["source_mime"] = media_body.mimetype()
            self.log.append(("update", fileId, f["name"]))
            return self._public(f)
        return _Req(run)


class FakeDriveService:
    def __init__(self):
        self.store: dict[str, dict] = {}
        self.log: list[tuple] = []
        self._files = FakeDriveFiles(self.store, self.log)

    def files(self):
        return self._files

    def by_name(self, name: str) -> dict:
        matches = [f for f in self.store.values() if f["name"] == name and not f["trashed"]]
        assert len(matches) == 1, f"{name}: {len(matches)} fichier(s)"
        return matches[0]

    def text(self, name: str) -> str:
        return self.by_name(name)["content"].decode("utf-8")


# --- Notion ----------------------------------------------------------------------------------------


class FakeNotion:
    """Imite `NotionClient.request` (et donc toutes les méthodes du vrai client qui s'appuient dessus)."""

    def __init__(self, fail_append_at: int | None = None):
        self.pages: dict[str, dict] = {}
        self.databases: dict[str, dict] = {}
        self.data_sources: dict[str, dict] = {}
        self.views: list[dict] = []
        self.calls: list[tuple] = []
        self._ids = itertools.count(1)
        self.fail_append_at = fail_append_at
        self.appends = 0
        self.root = "a" * 32
        self.pages[self.root] = {"id": self.root, "in_trash": False, "properties": {"title": {"type": "title", "title": [{"plain_text": "Cours M1"}]}}, "markdown": ""}

    def _id(self) -> str:
        return f"{next(self._ids):032x}"

    def request(self, method: str, path: str, *, json: dict | None = None, params: dict | None = None, max_attempts: int = 6) -> dict:
        from app.publish.notion import NotionError

        self.calls.append((method, path, json))
        parts = path.strip("/").split("/")
        if method == "GET" and parts == ["users", "me"]:
            return {"name": "Prise de notes"}
        if method == "POST" and parts == ["databases"]:
            db_id, ds_id = self._id(), self._id()
            props = {name: {"id": f"p{i}", **conf} for i, (name, conf) in enumerate(json["initial_data_source"]["properties"].items())}
            self.databases[db_id] = {"id": db_id, "in_trash": False, "data_sources": [{"id": ds_id}], "url": f"https://notion.example/{db_id}", "title": json["title"]}
            self.data_sources[ds_id] = {"id": ds_id, "properties": props}
            return self.databases[db_id]
        if method == "GET" and parts[0] == "databases":
            if parts[1] not in self.databases:
                raise NotionError(404, "object_not_found", "db")
            return self.databases[parts[1]]
        if parts[0] == "data_sources":
            ds = self.data_sources[parts[1]]
            if method == "PATCH":
                ds["properties"].update({k: {"id": f"x{k}", **v} for k, v in json["properties"].items()})
            return ds
        if method == "POST" and parts == ["views"]:
            self.views.append(json)
            return {"object": "view", "id": self._id()}
        if method == "POST" and parts == ["pages"]:
            pid = self._id()
            self.pages[pid] = {
                "id": pid, "in_trash": False, "parent": json["parent"], "properties": json["properties"],
                "markdown": json.get("markdown", ""), "url": f"https://notion.example/{pid}",
            }
            return self.pages[pid]
        if parts[0] == "pages" and len(parts) == 2:
            page = self.pages.get(parts[1])
            if page is None:
                raise NotionError(404, "object_not_found", "page")
            if method == "PATCH":
                if "properties" in json:
                    page["properties"].update(json["properties"])
                if "in_trash" in json:
                    page["in_trash"] = json["in_trash"]
            return page
        if parts[0] == "pages" and parts[2:] == ["markdown"]:
            page = self.pages[parts[1]]
            if method == "GET":
                return {"object": "page_markdown", "markdown": page["markdown"], "truncated": False, "unknown_block_ids": []}
            if json["type"] == "insert_content":
                self.appends += 1
                if self.fail_append_at is not None and self.appends == self.fail_append_at:
                    raise NotionError(502, "bad_gateway", "coupure simulée")
                page["markdown"] += "\n" + json["insert_content"]["content"]
            elif json["type"] == "replace_content":
                page["markdown"] = json["replace_content"]["new_str"]
            return {"object": "page_markdown", "markdown": page["markdown"]}
        raise AssertionError(f"Appel inattendu : {method} {path}")

    def session_pages(self) -> list[dict]:
        return [p for p in self.pages.values() if "Titre" in p.get("properties", {})]


def title_of(page: dict) -> str:
    prop = page["properties"]["Titre"]["title"]
    return "".join(t["text"]["content"] for t in prop)


def notion_like(md: str) -> str:
    """Approximation de ce que renvoie Notion en lecture (maths inline $`…`$)."""
    return re.sub(r"(?<!\$)\$([^$\n]+?)\$(?!\$)", r"$`\1`$", md)


# --- Documents de test (supports de cours) --------------------------------------------------------


def make_pdf(path, pages: list[list[str]]) -> bytes:
    """PDF minimal (Helvetica, une ligne de texte par élément) lisible par pypdf."""

    def esc(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    body: dict[int, bytes] = {1: b"<< /Type /Catalog /Pages 2 0 R >>",
                              3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"}
    page_ids = []
    for i, lines in enumerate(pages):
        pid, cid = 4 + 2 * i, 5 + 2 * i
        page_ids.append(pid)
        ops = ["BT", "/F1 18 Tf", "72 760 Td", "22 TL", *[f"({esc(line)}) Tj T*" for line in lines], "ET"]
        stream = "\n".join(ops).encode("cp1252")
        body[pid] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
                     f"/Contents {cid} 0 R >>").encode()
        body[cid] = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    body[2] = f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] /Count {len(pages)} >>".encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for i in sorted(body):
        offsets[i] = len(out)
        out += f"{i} 0 obj\n".encode() + body[i] + b"\nendobj\n"
    xref, size = len(out), max(body) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offsets[i]:010d} 00000 n \n".encode() for i in range(1, size))
    out += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return bytes(out)


def make_pptx(path, slides: list[tuple[str, list[str], str]]) -> bytes:
    """Présentation : (titre, puces, notes de l'intervenant) par diapositive."""
    from pptx import Presentation

    prs = Presentation()
    for title, bullets, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = bullets[0]
        for bullet in bullets[1:]:
            frame.add_paragraph().text = bullet
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))
    return path.read_bytes()
