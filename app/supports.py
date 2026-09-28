"""Supports de cours (diapositives, PDF, documents) joints à une séance.

Le fichier d'origine est conservé dans `recordings/<id>/supports/`. Son texte est lu par l'OCR de Mistral
(Markdown, formules en LaTeX, en-têtes et pieds de page retirés) ou, à défaut, extrait localement (pypdf,
python-pptx). Ce texte est transmis au LLM lors de la mise en forme : le cours suit ce qui a été dit, le
support sert à vérifier les termes, formules, notations et énoncés.
"""

from __future__ import annotations

import logging
import re
import threading
import uuid
from pathlib import Path
from typing import BinaryIO, Callable

from . import db, recorder
from .markdown_utils import convert_latex_delimiters, map_outside_code_math
from .retry import with_retries
from .textutils import slugify

log = logging.getLogger(__name__)

FORMATS = {".pdf": "PDF", ".pptx": "PowerPoint", ".docx": "Word"}
ACCEPT = ".pdf,.pptx,.docx,application/pdf,application/vnd.openxmlformats-officedocument.presentationml.presentation," \
         "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OFFICE = {".pptx", ".docx"}  # convertis (pas « vus ») par l'OCR : ponctuation échappée façon Markdown
MAGIC = {".pdf": b"%PDF", ".pptx": b"PK\x03\x04", ".docx": b"PK\x03\x04"}
MAX_BYTES = 50 * 1024 * 1024  # limite de l'OCR Mistral (50 Mo, 1 000 pages)
OCR_TIMEOUT_MS = 5 * 60 * 1000

_extract_lock = threading.Lock()  # une lecture à la fois (évite de lire deux fois le même fichier)


class SupportError(ValueError):
    pass


# --- Fichiers -------------------------------------------------------------------------------------

def supports_dir(recording_id: int) -> Path:
    return recorder.recording_dir(recording_id) / "supports"


def file_path(sup: dict) -> Path:
    return supports_dir(sup["recording_id"]) / sup["stored_name"]


def text_path(sup: dict) -> Path:
    return file_path(sup).with_suffix(".md")


def check_filename(filename: str) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in FORMATS:
        raise SupportError(f"« {filename} » : format non pris en charge (PDF, PPTX ou DOCX). "
                           "Pour Keynote, Google Slides ou un ancien .ppt, exportez d'abord en PDF.")
    return ext


def add(recording_id: int, filename: str, stream: BinaryIO) -> dict:
    """Enregistre le fichier (lu par morceaux, 50 Mo au plus) ; la lecture du texte se fait ensuite."""
    ext = check_filename(filename)
    folder = supports_dir(recording_id)
    folder.mkdir(parents=True, exist_ok=True)
    tmp = folder / f".envoi-{uuid.uuid4().hex}{ext}"
    size = 0
    head = b""
    try:
        with tmp.open("wb") as out:
            while chunk := stream.read(1 << 20):
                size += len(chunk)
                if size > MAX_BYTES:
                    raise SupportError(f"« {filename} » dépasse 50 Mo (limite de la lecture automatique).")
                if len(head) < 1024:
                    head += chunk[: 1024 - len(head)]
                out.write(chunk)
        if size == 0:
            raise SupportError(f"« {filename} » est vide.")
        if MAGIC[ext] not in head:
            raise SupportError(f"« {filename} » n'est pas un fichier {FORMATS[ext]} valide.")
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    sid = db.insert_support(recording_id=recording_id, filename=Path(filename).name, stored_name="", size=size)
    stored = f"{sid:03d}_{slugify(Path(filename).stem, 50)}{ext}"
    tmp.replace(folder / stored)
    db.update_support(sid, stored_name=stored)
    db.log(recording_id, f"Support de cours ajouté : « {Path(filename).name} ».")
    return db.get_support(sid)


def add_many(recording_id: int, files: list[tuple[str, BinaryIO]]) -> tuple[list[dict], list[str]]:
    """Ajoute plusieurs fichiers et lance leur lecture en arrière-plan ; renvoie (ajoutés, erreurs)."""
    added: list[dict] = []
    errors: list[str] = []
    for name, stream in files:
        if not name:
            continue
        try:
            sup = add(recording_id, name, stream)
        except SupportError as exc:
            errors.append(str(exc))
            continue
        extract_in_background(sup["id"])
        added.append(sup)
    return added, errors


def remove(support_id: int) -> None:
    sup = db.get_support(support_id)
    if not sup:
        return
    if sup["stored_name"]:
        file_path(sup).unlink(missing_ok=True)
        text_path(sup).unlink(missing_ok=True)
    db.delete_support(support_id)
    db.log(sup["recording_id"], f"Support de cours retiré : « {sup['filename']} ».")


# --- Lecture du texte -----------------------------------------------------------------------------

_IMAGE_RE = re.compile(r"!\[[^\]]*\]\([^)\s]*\)")
_TABLE_REF_RE = re.compile(r"\[(tbl-\d+\.(?:md|html))\]\(\1\)")
_ESCAPED_RE = re.compile(r"\\([!-/:-@\[-`{-~])")  # ponctuation ASCII échappée (\( \, \+ \_ …)


def clean_ocr_markdown(md: str, *, office: bool, tables: list | None = None) -> str:
    """Nettoie une page renvoyée par l'OCR : figures, tableaux séparés, échappements, délimiteurs LaTeX.

    Pour un PDF (lu comme une image), `\\( … \\)` délimite une formule : on la passe en `$…$` comme dans le
    cours. Pour un PPTX/DOCX (converti), `\\(` n'est qu'une parenthèse échappée : on retire l'échappement.
    """
    contents = {getattr(t, "id", None): getattr(t, "content", "") for t in tables or []}
    md = _TABLE_REF_RE.sub(lambda m: "\n" + (contents.get(m.group(1)) or "") + "\n", md or "")
    md = _IMAGE_RE.sub("[figure]", md)
    if office:
        md = map_outside_code_math(md, lambda seg: _ESCAPED_RE.sub(r"\1", seg))
    else:
        md = convert_latex_delimiters(md)
    md = re.sub(r"[ \t]+\n", "\n", md)
    return re.sub(r"\n{3,}", "\n\n", md).strip()


def ocr_pages(path: Path) -> list[str]:
    """Texte de chaque page (ou diapositive) par l'OCR Mistral ; le fichier envoyé est supprimé ensuite."""
    from .llm import get_client

    client = get_client()
    model = db.get_setting("ocr_model").strip() or "mistral-ocr-latest"
    data = path.read_bytes()
    uploaded = with_retries(
        lambda: client.files.upload(file={"file_name": path.name, "content": data}, purpose="ocr"),
        attempts=4, label="Envoi du support à l'OCR",
    )
    try:
        res = with_retries(
            lambda: client.ocr.process(
                model=model,
                document={"type": "file", "file_id": uploaded.id},
                extract_header=True,
                extract_footer=True,
                timeout_ms=OCR_TIMEOUT_MS,
            ),
            attempts=4, label="OCR du support",
        )
    finally:
        try:
            client.files.delete(file_id=uploaded.id)
        except Exception as exc:  # noqa: BLE001 - le fichier expirera de toute façon côté Mistral
            log.warning("Fichier OCR %s non supprimé : %s", uploaded.id, exc)
    office = path.suffix.lower() in OFFICE
    return [clean_ocr_markdown(p.markdown, office=office, tables=p.tables) for p in sorted(res.pages, key=lambda p: p.index)]


def _shape_texts(shapes) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            out += _shape_texts(shape.shapes)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                out.append(" | ".join(cell.text.strip() for cell in row.cells))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
            out.append(shape.text_frame.text.strip())
    return out


def _pptx_notes(path: Path) -> list[str]:
    """Notes de l'intervenant de chaque diapositive (souvent riches, jamais lues par l'OCR)."""
    from pptx import Presentation

    notes = []
    for slide in Presentation(str(path)).slides:
        text = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        notes.append(text)
    return notes


def local_pages(path: Path) -> list[str]:
    """Extraction locale (sans réseau), moins fidèle que l'OCR : pas de formules ni de mise en page."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        return [(page.extract_text() or "").strip() for page in PdfReader(str(path)).pages]
    if ext == ".pptx":
        from pptx import Presentation

        return ["\n".join(_shape_texts(slide.shapes)) for slide in Presentation(str(path)).slides]
    raise SupportError("pas de lecture locale pour ce format")


def _with_notes(pages: list[str], notes: list[str]) -> list[str]:
    if len(notes) != len(pages):
        return pages
    return [f"{p}\n\nNotes de l'intervenant : {n}" if n else p for p, n in zip(pages, notes)]


def render_pages(pages: list[str], ext: str) -> str:
    label = "Diapositive" if ext == ".pptx" else "Page"
    return "\n\n".join(f"[{label} {i}]\n{text.strip() or '(vide)'}" for i, text in enumerate(pages, 1)) + "\n"


def extract(support_id: int) -> dict | None:
    """Lit le texte du support : OCR Mistral, sinon extraction locale. Sans effet s'il est déjà lu."""
    with _extract_lock:
        sup = db.get_support(support_id)
        if sup is None or not sup["stored_name"]:
            return sup
        if sup["status"] == "done" and text_path(sup).exists():
            return sup
        db.update_support(support_id, status="extracting", error=None)
        path = file_path(sup)
        ext = path.suffix.lower()
        errors: list[str] = []
        pages: list[str] = []
        method = None
        try:
            pages, method = ocr_pages(path), "ocr"
        except Exception as exc:  # noqa: BLE001 - on tente la lecture locale
            log.warning("OCR du support %s impossible : %s", support_id, exc)
            errors.append(f"OCR : {exc}")
        if not any(p.strip() for p in pages) and ext in (".pdf", ".pptx"):
            try:
                pages, method = local_pages(path), "local"
            except Exception as exc:  # noqa: BLE001
                errors.append(f"lecture locale : {exc}")
        if ext == ".pptx" and any(p.strip() for p in pages):
            try:
                pages = _with_notes(pages, _pptx_notes(path))
            except Exception as exc:  # noqa: BLE001 - les notes sont un plus
                log.warning("Notes du support %s illisibles : %s", support_id, exc)
        if not any(p.strip() for p in pages):
            message = " ; ".join(errors) or "aucun texte trouvé (document scanné illisible ou vide ?)"
            db.update_support(support_id, status="error", error=message[:1000], method=None)
            db.log(sup["recording_id"], f"Support « {sup['filename']} » illisible : {message}", "warning")
            return db.get_support(support_id)
        if db.get_support(support_id) is None:  # retiré pendant la lecture
            return None
        text = render_pages(pages, ext)
        text_path(sup).write_text(text, encoding="utf-8")
        db.update_support(support_id, status="done", error=None, method=method, pages=len(pages), chars=len(text))
        how = "OCR Mistral" if method == "ocr" else "lecture locale"
        db.log(sup["recording_id"], f"Support « {sup['filename']} » lu : {len(pages)} page(s), {how}.")
        return db.get_support(support_id)


def _extract_safely(support_id: int) -> None:
    try:
        extract(support_id)
    except Exception as exc:  # noqa: BLE001 - un fil d'arrière-plan ne doit rien laisser « en cours »
        log.exception("Lecture du support %s en échec", support_id)
        db.update_support(support_id, status="error", error=str(exc)[:1000])


def extract_in_background(support_id: int) -> None:
    threading.Thread(target=_extract_safely, args=(support_id,), name=f"support-{support_id}", daemon=True).start()


def resume_pending() -> None:
    """Au démarrage : relit les supports dont la lecture a été interrompue (arrêt de l'app)."""
    for sup in db.list_supports_by_status("pending", "extracting"):
        extract_in_background(sup["id"])


# --- Pour la mise en forme ------------------------------------------------------------------------

def prepare(recording_id: int, on_progress: Callable[[str], None] | None = None) -> list[dict]:
    """Lit les supports pas encore lus (ou en échec : nouvel essai) ; renvoie ceux qui sont lisibles."""
    for sup in db.list_supports(recording_id):
        if sup["status"] != "done" or not text_path(sup).exists():
            if on_progress:
                on_progress(f"Lecture du support « {sup['filename']} »…")
            extract(sup["id"])
    return [s for s in db.list_supports(recording_id) if s["status"] == "done" and text_path(s).exists()]


_PAGE_RE = re.compile(r"^\[(?:Page|Diapositive) \d+\]$", re.M)


def split_pages(text: str) -> list[tuple[str, str]]:
    """[(repère, texte)] depuis le texte enregistré (« [Page N] » / « [Diapositive N] » en tête de page)."""
    labels = [m.group(0)[1:-1] for m in _PAGE_RE.finditer(text)]
    bodies = _PAGE_RE.split(text)[1:]
    return [(label, body.strip()) for label, body in zip(labels, bodies)]


def documents(sups: list[dict]) -> list[dict]:
    """Supports lus, découpés en pages, pour la mise en forme (voir llm.format_course)."""
    return [{"name": s["filename"], "pages": split_pages(text_path(s).read_text(encoding="utf-8"))} for s in sups]


def mark_used(support_ids: list[int]) -> None:
    now = db.now_iso()
    for sid in support_ids:
        db.update_support(sid, used_at=now)


def summary(recording_id: int) -> dict:
    """État des supports d'une séance pour l'interface."""
    sups = db.list_supports(recording_id)
    return {
        "supports": sups,
        "reading": any(s["status"] in ("pending", "extracting") for s in sups),
        "unused": [s for s in sups if s["status"] == "done" and not s["used_at"]],
    }
